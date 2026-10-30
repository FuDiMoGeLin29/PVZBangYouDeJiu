
import openpyxl, json
p = r"D:\Microsoft Visual Studio\Code\PVZBangYouDeJiu\PVZBangYouDeJiu\Docs\pvz1051函数表3.3743.1157.xlsx"
wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
print("SHEETS:", wb.sheetnames)
for ws in wb.worksheets:
    print("---", ws.title, ws.max_row, ws.max_column)
